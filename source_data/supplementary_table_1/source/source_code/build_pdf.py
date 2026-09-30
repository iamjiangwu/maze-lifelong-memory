"""Render the current source-backed S1 workbook as a one-page publication table."""
import argparse
import hashlib
import json
from pathlib import Path
from xml.sax.saxutils import escape

import fitz
from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--output-dir', type=Path, default=ROOT / 'build')
    args = parser.parse_args()
    output = args.output_dir.resolve()
    source = json.loads((ROOT / 'source_data/current_rows.json').read_text(encoding='utf-8'))
    source_xlsx = output / 'Supplementary_Table_S1_Mouse_Cohorts_by_Analysis.xlsx'
    workbook = load_workbook(source_xlsx, data_only=True)
    sheet = workbook['Table S1']
    count = len(source['rows'])
    merges = [m for m in sheet.merged_cells.ranges if 4 <= m.min_row <= m.max_row <= count + 3]
    rows = [[sheet.cell(r, c).value for c in range(1, 5)] for r in range(4, count + 4)]
    checks = []
    for i, src in enumerate(source['rows']):
        expected = [src['analysis'], src['panels'], f"n = {len(src['mice'])}\n" + ', '.join(src['mice']), src['note']]
        for c, value in enumerate(expected, 1):
            cell = sheet.cell(i+4, c)
            actual = cell.value
            for merged in merges:
                if cell.coordinate in merged:
                    actual = sheet.cell(merged.min_row, merged.min_col).value
            checks.append(actual == value or (actual is None and value == ''))
    assert all(checks), 'Workbook values differ from current_rows.json'
    workbook.close()
    page = landscape(A4)
    pdf = output / 'Supplementary_Table_S1_Mouse_Cohorts_by_Analysis.pdf'
    doc = SimpleDocTemplate(str(pdf), pagesize=page, leftMargin=21, rightMargin=21, topMargin=23, bottomMargin=23, title=source['title'], author='', invariant=1)
    title_style = ParagraphStyle('Title', fontName='Helvetica-Bold', fontSize=12.2, leading=14.2)
    body = ParagraphStyle('Body', fontName='Helvetica', fontSize=6.8, leading=7.9, splitLongWords=False, textColor=colors.HexColor('#1A1A1A'))
    header = ParagraphStyle('Header', parent=body, fontName='Helvetica-Bold', fontSize=7.2, leading=8.2)
    mice = ParagraphStyle('Mice', parent=body)
    group = ParagraphStyle('Group', parent=body, fontName='Helvetica-Bold')
    def paragraph(value, style, bold_first=False):
        lines = escape('' if value is None else str(value)).splitlines()
        if bold_first and lines:
            lines[0] = '<b>' + lines[0] + '</b>'
        return Paragraph('<br/>'.join(lines), style)
    data = [[paragraph(v, header) for v in source['headers']]]
    data += [[paragraph(v, [body, body, mice, body][c], c == 2) for c, v in enumerate(row)] for row in rows]
    available = page[0] - 42
    table = Table(data, colWidths=[152, 123, 172, available - 447], hAlign='LEFT', repeatRows=1)
    style = TableStyle([
        ('BACKGROUND', (0,0),(-1,0), colors.HexColor('#E4E4E4')),
        ('VALIGN',(0,0),(-1,-1),'TOP'),
        ('LEFTPADDING',(0,0),(-1,-1),3),('RIGHTPADDING',(0,0),(-1,-1),3),
        ('TOPPADDING',(0,0),(-1,-1),3),('BOTTOMPADDING',(0,0),(-1,-1),3),
        ('GRID',(0,0),(-1,-1),0.22,colors.HexColor('#C8C8C8')),
        ('BOX',(0,0),(-1,-1),0.55,colors.HexColor('#444444')),
        ('LINEBELOW',(0,0),(-1,0),0.55,colors.HexColor('#444444')),
    ])
    for merged in merges:
        start=(merged.min_col-1, merged.min_row-3)
        end=(merged.max_col-1, merged.max_row-3)
        style.add('SPAN',start,end)
        if merged.min_col == 3:
            style.add('VALIGN',start,end,'MIDDLE')
    table.setStyle(style)
    width, height = table.wrap(available, page[1]-70)
    assert height <= page[1]-85, f'Table too tall: {height:.1f} pt'
    doc.build([paragraph(source['title'],title_style),Spacer(1,6),table])
    document=fitz.open(pdf)
    assert len(document)==1, 'S1 must remain one page'
    document[0].get_pixmap(matrix=fitz.Matrix(2.5,2.5)).save(output/'Supplementary_Table_S1_publication.png')
    text=document[0].get_text()
    (output/'publication_text.txt').write_text(text,encoding='utf-8')
    validation={'status':'PASS_PENDING_VISUAL_QA','rows':count,'source_cell_checks':len(checks),'pages':1,'table_height_pt':height,'body_font_pt':6.8,'pdf_sha256':hashlib.sha256(pdf.read_bytes()).hexdigest(),'workbook_sha256':hashlib.sha256(source_xlsx.read_bytes()).hexdigest()}
    (output/'PDF_BUILD_VALIDATION.json').write_text(json.dumps(validation,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(validation,indent=2))


if __name__ == '__main__':
    main()
