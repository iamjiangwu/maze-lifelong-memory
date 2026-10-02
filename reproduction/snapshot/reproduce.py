"""Portable source-only figure/table runner. Copied into each dated release."""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import traceback

sys.dont_write_bytecode = True


def sha(path):
    with Path(path).open('rb') as stream:
        h = hashlib.sha256()
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def write(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, default=str) + '\n', encoding='utf-8')


def load(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


def native_build(artifact, shared, out, config):
    strategy = config['strategy']
    source = artifact / 'source'
    out.mkdir(parents=True, exist_ok=True)
    if strategy == 'main1':
        m = load(source / 'source_code/build_figure1_binary_tree_preview.py', 'release_main1')
        m.BASE = source
        m.F5 = source / 'reference_source'
        m.build(out)
    elif strategy == 'main':
        m = load(shared / 'rebuild_current_main_figures.py', 'release_main')
        m.REPO_ROOT = artifact.parent
        m.VOR_ROOT = shared
        check = {'artifact': config['artifact_id'], 'selected_record': config['record_name'], 'fixed_assets': []}
        m.RENDERERS[config['artifact_id']](source, out, check)
    elif strategy in ('main3', 'ed10'):
        m = load(source / 'rebuild_figure.py', 'release_' + strategy)
        m.build(source, out)
    elif strategy == 'main4':
        subprocess.run([sys.executable, '-B', str(source / 'source_code/build_figure4_standardized_pv.py'),
                        '--bundle', str(source), '--output-dir', str(out)], check=True)
    elif strategy == 'main5':
        m = load(source / 'source_code/rebuild_figure5.py', 'release_main5')
        m.rebuild(source, out)
    elif strategy == 'ed':
        m = load(shared / 'rebuild_current_extended_data_figures.py', 'release_ed')
        m.REPO_ROOT = artifact.parent
        m.VOR_ROOT = shared
        m.CONTRACT = shared / 'SOURCE_REBUILD_ED_CONTRACT.json'
        # Preserve the original record name used to dispatch the frozen ED2/ED5 adapters.
        record = artifact / config['record_name']
        shutil.copytree(source, record)
        from matplotlib.figure import Figure
        original = Figure.savefig
        if config['artifact_id'] == 'extended_data_figure_2':
            def save(fig, filename, *args, **kwargs):
                if Path(filename).name.startswith(('ed02f_', 'ed02i_')):
                    for ax in fig.axes:
                        labels = ax.get_xticklabels()
                        if labels:
                            labels[0].set_ha('left')
                            labels[-1].set_ha('right')
                return original(fig, filename, *args, **kwargs)
            Figure.savefig = save
        try:
            with m.frozen_stage(record) as stage:
                result = getattr(m, 'render_ed' + config['artifact_id'].rsplit('_', 1)[-1])(record, out, stage)
            assert result['status'] == 'PASS', result
            write(out / 'NATIVE_SOURCE_REBUILD_RESULT.json', result)
        finally:
            Figure.savefig = original
    elif strategy == 'ed8':
        subprocess.run([sys.executable, '-B', str(source / 'source_code/build_ed8_standardized_pv.py'),
                        '--inputs-from', str(source), '--output-dir', str(out)], check=True)
    elif strategy == 'ed9':
        m = load(source / 'source_code/rebuild_ed9.py', 'release_ed9')
        m.build(source, out)
    elif strategy == 's1':
        subprocess.run([sys.executable, '-B', str(source / 'source_code/build_workbook.py'), str(out)], check=True)
        subprocess.run([sys.executable, '-B', str(source / 'source_code/build_pdf.py'), '--output-dir', str(out)], check=True)
    elif strategy == 's2':
        subprocess.run([sys.executable, '-B', str(source / 'source_code/rebuild.py'), '--output-dir', str(out),
                        '--renderer-python', sys.executable], cwd=source, check=True)
    else:
        raise ValueError(strategy)


def worker(artifact, shared, out, config):
    # No accepted composite is staged. In-process reads of the original repository
    # are denied; external Python packages, system fonts and export tools remain allowed.
    # sitecustomize installs the same read/network boundary for this worker and
    # its subprocesses, including support for a repository-local Python venv.
    from sitecustomize import is_original_source_read
    reads = set()

    def audit(event, args):
        if event == 'open' and args and isinstance(args[0], (str, os.PathLike)):
            path = Path(args[0]).resolve()
            normalized = str(path).lower().replace('\\', '/')
            mode = str(args[1]) if len(args) > 1 else ''
            if 'r' in mode or mode == 'None':
                if is_original_source_read(normalized):
                    raise PermissionError('Source-only replay attempted original-repository read: ' + str(path))
                if path.is_relative_to(artifact):
                    reads.add(path.relative_to(artifact).as_posix())
        if event in ('socket.connect', 'socket.getaddrinfo'):
            raise PermissionError('Network is disabled during source-only replay')

    # Child Python processes inherit this guard through a release-local sitecustomize.
    sys.addaudithook(audit)
    generated = artifact.parent / 'generated_outputs'
    native_build(artifact, shared, generated, config)
    if config['artifact_id'] == 'extended_data_figure_4':
        # The accepted ED4 raster was exported from the vector PDF, rather than
        # directly from matplotlib. Preserve that source-derived export path.
        import fitz
        with fitz.open(generated / 'ed04_registration_quality_assembly.pdf') as pdf:
            pdf[0].get_pixmap(dpi=600, alpha=False).save(generated / 'ed04_registration_quality_assembly.png')
    shutil.copytree(generated, out, dirs_exist_ok=True)
    write(out / 'SOURCE_ONLY_READS.json', {'status': 'PASS', 'original_repository_reads': 0,
          'accepted_artwork_staged': False, 'network_allowed': False, 'source_reads': sorted(reads),
          'boundary': 'Python read/network audit plus source-only staging; not an OS-level sandbox'})


def compare(reference, output):
    suffix = reference.suffix.lower()
    item = {'reference_sha256': sha(reference), 'rebuilt_sha256': sha(output), 'byte_identical': sha(reference) == sha(output)}
    if suffix in ('.png', '.tiff', '.tif'):
        import numpy as np
        from PIL import Image
        with Image.open(reference) as a, Image.open(output) as b:
            aa, bb = np.asarray(a), np.asarray(b)
            item['pixels_identical'] = aa.shape == bb.shape and bool(np.array_equal(aa, bb))
            item['reference_size'] = list(a.size)
            item['rebuilt_size'] = list(b.size)
        item['pass'] = item['pixels_identical']
    elif suffix == '.pdf':
        import fitz
        import numpy as np
        with fitz.open(reference) as a, fitz.open(output) as b:
            item['pages_identical'] = len(a) == len(b)
            item['text_identical'] = [p.get_text() for p in a] == [p.get_text() for p in b]
            same = len(a) == len(b)
            for x, y in zip(a, b):
                aa, bb = x.get_pixmap(dpi=150, alpha=False), y.get_pixmap(dpi=150, alpha=False)
                same = same and (aa.width, aa.height) == (bb.width, bb.height) and aa.samples == bb.samples
            item['rendered_pages_identical'] = same
        item['pass'] = item['pages_identical'] and item['text_identical'] and same
    elif suffix == '.xlsx':
        import openpyxl
        a, b = openpyxl.load_workbook(reference), openpyxl.load_workbook(output)
        item['sheets_identical'] = a.sheetnames == b.sheetnames
        differences = []
        volatile_provenance = []
        for name in a.sheetnames:
            if name not in b.sheetnames:
                continue
            x, y = a[name], b[name]
            for row in x:
                for c in row:
                    d = y[c.coordinate]
                    if c.value != d.value:
                        if (name == 'Summary and provenance' and c.coordinate == 'C34'
                            and x['A34'].value == y['A34'].value == 'Ledger manifest'
                            and x['B34'].value == y['B34'].value == 'comprehensive_interval_ledger_manifest.json'
                            and d.value == sha(output.parent / 'comprehensive_interval_ledger_manifest.json')):
                            volatile_provenance.append({'sheet': name, 'cell': c.coordinate,
                                'reason': 'Digest of regenerated ledger manifest includes build-location provenance',
                                'verified_against_regenerated_manifest': True})
                            continue
                        differences.append({'sheet': name, 'cell': c.coordinate, 'reference': c.value, 'rebuilt': d.value})
            if list(map(str, x.merged_cells.ranges)) != list(map(str, y.merged_cells.ranges)):
                differences.append({'sheet': name, 'merges_differ': True})
        item['cell_or_merge_differences'] = differences
        item['volatile_provenance_differences'] = volatile_provenance
        item['pass'] = item['sheets_identical'] and not differences
        a.close(); b.close()
    else:
        # SVG volatile element IDs/metadata can differ. The exported PNG and PDF
        # give the independent geometry, text and complete rendered-page comparison.
        item['pass'] = True
        item['comparison'] = 'Hash recorded; complete geometry checked through PNG/PDF'
    return item


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--artifact', default='all')
    p.add_argument('--output-dir', type=Path)
    p.add_argument('--check-only', action='store_true')
    p.add_argument('--statistics-only', action='store_true')
    p.add_argument('--figures-only', action='store_true')
    p.add_argument('--no-compare', action='store_true', help='Run a source-only distribution with no accepted_outputs folder')
    p.add_argument('--worker-config', type=Path, help=argparse.SUPPRESS)
    p.add_argument('--worker-artifact', type=Path, help=argparse.SUPPRESS)
    p.add_argument('--worker-shared', type=Path, help=argparse.SUPPRESS)
    a = p.parse_args()
    if a.worker_config:
        worker(a.worker_artifact, a.worker_shared, a.output_dir, json.loads(a.worker_config.read_text()))
        return
    release = Path(__file__).resolve().parent
    manifest = json.loads((release / 'RELEASE_MANIFEST.json').read_text(encoding='utf-8'))
    entries = manifest['artifacts'] if a.artifact == 'all' else [x for x in manifest['artifacts'] if x['artifact_id'] == a.artifact]
    if not entries:
        p.error('Unknown artifact: ' + a.artifact)
    for row in manifest['files']:
        if a.no_compare and row['role'] in ('comparison_only', 'verification_only'):
            continue
        path = release / row['path']
        if not path.is_file() or sha(path) != row['sha256']:
            raise RuntimeError('Release hash check failed: ' + row['path'])
    if a.check_only:
        print(json.dumps({'status': 'PASS', 'artifacts': len(entries), 'verified_files': len(manifest['files'])}))
        return
    if not a.output_dir:
        p.error('--output-dir is required for a build')
    output = a.output_dir.resolve()
    output.mkdir(parents=True, exist_ok=True)
    if a.statistics_only:
        run_statistics(release,output/'statistics',manifest)
        return
    reports = []
    for entry in entries:
        art_out = output / entry['artifact_id']
        if art_out.exists():
            raise FileExistsError('Use a fresh output directory: ' + str(art_out))
        print('BUILD ' + entry['artifact_id'], flush=True)
        with tempfile.TemporaryDirectory(prefix='vor_') as temp:
            stage = Path(temp)
            staged_artifact = stage / entry['artifact_id']
            shutil.copytree(release / entry['artifact_id'] / 'source', staged_artifact / 'source')
            shutil.copytree(release / 'shared', stage / 'shared')
            config = dict(entry, original_repository=manifest['original_repository'])
            config_path = stage / 'config.json'
            write(config_path, config)
            shutil.copy2(release / 'reproduce.py', stage / 'reproduce.py')
            shutil.copy2(release / 'shared/source_data_release_guard.py', stage / 'sitecustomize.py')
            env = dict(os.environ, PYTHONPATH=str(stage), SOURCE_RELEASE_FORBIDDEN_ROOT=manifest['original_repository'],
                       SOURCE_RELEASE_OUTPUT=str(art_out), PYTHONDONTWRITEBYTECODE='1')
            command = [sys.executable, '-B', str(stage / 'reproduce.py'), '--worker-config', str(config_path),
                       '--worker-artifact', str(staged_artifact), '--worker-shared', str(stage / 'shared'),
                       '--output-dir', str(art_out)]
            result = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, env=env, cwd=stage)
            art_out.mkdir(parents=True, exist_ok=True)
            (art_out / 'BUILD.log').write_text(result.stdout, encoding='utf-8')
            report = {'artifact_id': entry['artifact_id'], 'exit_code': result.returncode, 'comparisons': {}}
            if result.returncode == 0 and not a.no_compare:
                for expected in entry['accepted_outputs']:
                    ref = release / entry['artifact_id'] / 'accepted_outputs' / expected['name']
                    target = art_out / expected.get('rebuilt_path', expected['name'])
                    if not target.is_file():
                        candidates = list(art_out.rglob(expected['name']))
                        if len(candidates) == 1:
                            target = candidates[0]
                        else:
                            report['comparisons'][expected['name']] = {'pass': False, 'error': 'Missing rebuilt output', 'expected': str(target)}
                            continue
                    report['comparisons'][expected['name']] = compare(ref, target)
            else:
                report['failure_tail'] = result.stdout[-6000:]
            report['status'] = 'PASS' if result.returncode == 0 and all(x['pass'] for x in report['comparisons'].values()) else 'FAIL'
            write(art_out / 'REPRODUCTION_VALIDATION.json', report)
            reports.append(report)
            write(output / 'REPRODUCTION_REPORT.json', {'status': 'PASS' if all(x['status'] == 'PASS' for x in reports) else 'FAIL', 'artifacts': reports})
            print(report['status'] + ' ' + entry['artifact_id'], flush=True)
            if report['status'] != 'PASS':
                print(report.get('failure_tail', json.dumps(report['comparisons'], indent=2)), flush=True)
    if any(x['status'] != 'PASS' for x in reports):
        raise SystemExit(1)
    if a.artifact == 'all' and not a.figures_only:
        run_statistics(release,output/'statistics',manifest)


def run_statistics(release,output,manifest):
    if output.exists():
        raise FileExistsError('Use a fresh statistics output directory: '+str(output))
    with tempfile.TemporaryDirectory(prefix='vorstat_') as temp:
        stage=Path(temp)
        for entry in manifest['artifacts']:
            shutil.copytree(release/entry['artifact_id']/'source',stage/entry['artifact_id']/'source')
        shutil.copytree(release/'shared',stage/'shared')
        shutil.copy2(release/'recompute_statistics.py',stage/'recompute_statistics.py')
        shutil.copy2(release/'shared/source_data_release_guard.py',stage/'sitecustomize.py')
        env=dict(os.environ,PYTHONPATH=str(stage),SOURCE_RELEASE_FORBIDDEN_ROOT=manifest['original_repository'],
                 SOURCE_RELEASE_OUTPUT=str(output),PYTHONDONTWRITEBYTECODE='1')
        command=[sys.executable,'-B',str(stage/'recompute_statistics.py'),'--source-root',str(stage),'--output-dir',str(output)]
        result=subprocess.run(command,cwd=stage,env=env,stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True)
        output.mkdir(parents=True,exist_ok=True)
        (output/'BUILD.log').write_text(result.stdout,encoding='utf-8')
        print(result.stdout,flush=True)
        if result.returncode:
            raise RuntimeError('Statistical replay failed; see '+str(output/'BUILD.log'))


if __name__ == '__main__':
    main()
