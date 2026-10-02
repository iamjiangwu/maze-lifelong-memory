"""Build Supplementary Table 2 with openpyxl and source-derived formula caches.

openpyxl writes formulas but does not calculate them. The specific formula
results below are derived independently from the same source ledger and cached
in the XLSX. Excel recalculates on opening. No accepted workbook is a build input.
"""
import argparse
import csv
import hashlib
import io
import json
import statistics
import xml.etree.ElementTree as ET
from copy import copy
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.properties import CalcProperties

SHORT = '1616day1->1616bday1'
DOUBLE = '618-6day1->618-12day1'
RETAINED = '1616bday1->1616b-3day1'
SENTENCE = "Day~$i$ of the first cycle was the animal's naive session; subsequent Day~$i$ sessions were generally return sessions conducted after approximately 90~d without maze experience. Across 56 included long-gap intervals, the mean $\\pm$ s.d. was 90 $\\pm$ 4~d (range, 82--110~d). Mouse 1616's 14-d interval and mouse 618's 180-d interval were excluded from this calculation; actual interval lengths are summarized in Supplementary Table~2."
INK, NAVY, GRAY, PALE = '202020', '1F4168', 'E4E4E4', 'F7F8FA'
THIN, DARK = Side(style='thin', color='FFC8C8C8'), Side(style='thin', color='FF555555')


def style(sheet, address, *, size=None, color=INK, bold=False, fill=None,
          horizontal=None, vertical=None, wrap=None, number_format=None):
    for row in sheet[address]:
        for cell in row:
            if size is not None: cell.font = Font(name='Carlito', size=size, color='FF' + color, bold=bold, italic=False)
            if fill is not None: cell.fill = PatternFill('solid', fgColor='FF' + fill)
            alignment = copy(cell.alignment)
            if horizontal is not None: alignment.horizontal = horizontal
            if vertical is not None: alignment.vertical = vertical
            if wrap is not None: alignment.wrap_text = wrap
            cell.alignment = alignment
            if number_format is not None: cell.number_format = number_format


def values(sheet, start, rows):
    for r, row in enumerate(rows, start):
        for c, value in enumerate(row, 1): sheet.cell(r, c, value)


def previous_month(value):
    if not value: return 0
    try: return float(value)
    except ValueError: return None


def cache_formulas(path, caches):
    namespace = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    ET.register_namespace('', namespace)
    with ZipFile(path) as archive: entries = [(item, archive.read(item.filename)) for item in archive.infolist()]
    result = io.BytesIO()
    with ZipFile(result, 'w') as archive:
        for info, payload in entries:
            if info.filename in caches:
                root = ET.fromstring(payload)
                seen = set()
                for cell in root.iter(f'{{{namespace}}}c'):
                    if cell.find(f'{{{namespace}}}f') is None: continue
                    coordinate = cell.attrib['r']
                    value = caches[info.filename][coordinate]
                    seen.add(coordinate)
                    node = cell.find(f'{{{namespace}}}v')
                    if node is None: node = ET.SubElement(cell, f'{{{namespace}}}v')
                    if isinstance(value, str): cell.set('t', 'str'); node.text = value
                    else: cell.attrib.pop('t', None); node.text = repr(value)
                assert seen == set(caches[info.filename]), 'Formula cache address mismatch'
                payload = ET.tostring(root, encoding='utf-8', xml_declaration=True)
            archive.writestr(info, payload)
    path.write_bytes(result.getvalue())


def build(output, source, notes_source):
    source_path = source / 'comprehensive_interval_ledger.csv'
    manifest_path = source / 'comprehensive_interval_ledger_manifest.json'
    text = source_path.read_bytes().decode('utf-8')
    rows = list(csv.DictReader(io.StringIO(text)))
    manifest_text = manifest_path.read_bytes().decode('utf-8')
    manifest = json.loads(manifest_text)
    notes = json.loads((notes_source / 'publication_notes.json').read_text(encoding='utf-8'))
    yes = lambda r, key: r[key].lower() == 'true'
    animals = sorted({r['Animal ID'] for r in rows}, key=int)
    behavior = [r for r in rows if yes(r, 'Fixed-month behavior return roster')]
    calcium = [r for r in rows if yes(r, 'Maintained calcium imaging')]
    overlap = [r for r in behavior if yes(r, 'Maintained calcium imaging')]
    included = [r for r in rows if r['Pair key'] not in [SHORT, DOUBLE]]
    days = [int(r['Interval (d)']) for r in included]
    assert [len(rows), len(animals), len(behavior), len(calcium), len(overlap), len(days)] == [58, 18, 56, 37, 35, 56]
    assert abs(statistics.mean(days) - 90.08928571428571) < 1e-10
    assert abs(statistics.stdev(days) - 3.56384259276052) < 1e-10
    assert next(r for r in rows if r['Pair key'] == RETAINED)['Return month'] == '6'
    assert next(r for r in rows if r['Pair key'] == DOUBLE)['Return month'] == '12'
    book = Workbook()
    book._fonts[0] = Font(name='Carlito', size=11)  # Original workbook's default style.
    book.remove(book.active)
    pub, ledger, summary = [book.create_sheet(name) for name in ['Supplementary Table 2', 'Source Data', 'Summary and provenance']]
    for sheet in book: sheet.sheet_view.showGridLines = False
    book.calculation = CalcProperties(fullCalcOnLoad=True, forceFullCalc=True)
    caches = {f'xl/worksheets/sheet{i}.xml': {} for i in range(1, 4)}
    pc, lc, sc = list(caches.values())
    pub.merge_cells('A1:G1')
    pub['A1'] = 'Supplementary Table 2. Intervals between successive recording cycles'
    style(pub, 'A1:G1', size=17, bold=True, vertical='center')
    pub.row_dimensions[1].height = 34
    values(pub, 3, [['Mouse', '3-month return\n(d)', '6-month return\n(d)', '9-month return\n(d)', '12-month return\n(d)', 'Other intervals\n(d)', 'Calcium-imaged returns\n(months)']])
    style(pub, 'A3:G3', size=10, bold=True, fill=GRAY, horizontal='center', vertical='center', wrap=True)
    for c in range(1, 8): pub.cell(3, c).border = Border(top=DARK, bottom=DARK, left=THIN if c > 1 else Side(), right=THIN if c < 7 else Side())
    pub.row_dimensions[3].height = 42
    short_excel = next(i for i, r in enumerate(rows, 2) if r['Pair key'] == SHORT)
    for r, animal in enumerate(animals, 4):
        animal_rows = [x for x in rows if x['Animal ID'] == animal]
        pub.cell(r, 1, animal)
        for column, month in [('B', 3), ('C', 6), ('D', 9), ('E', 12)]:
            coord = f'{column}{r}'
            pub[coord] = f'=IFERROR(SUMIFS(\'Source Data\'!$I$2:$I$59,\'Source Data\'!$A$2:$A$59,$A{r},\'Source Data\'!$B$2:$B$59,{month})/COUNTIFS(\'Source Data\'!$A$2:$A$59,$A{r},\'Source Data\'!$B$2:$B$59,{month}),"")'
            selected = [int(x['Interval (d)']) for x in animal_rows if x['Return month'] == str(month)]
            pc[coord] = statistics.mean(selected) if selected else ''
        coord = f'F{r}'
        if animal == '1616': pub[coord] = f"='Source Data'!I{short_excel}"; pc[coord] = 14
        else:
            pub[coord] = f'=IFERROR(SUMIFS(\'Source Data\'!$I$2:$I$59,\'Source Data\'!$A$2:$A$59,$A{r},\'Source Data\'!$J$2:$J$59,"No")/COUNTIFS(\'Source Data\'!$A$2:$A$59,$A{r},\'Source Data\'!$J$2:$J$59,"No"),"")'
            selected = [int(x['Interval (d)']) for x in animal_rows if not yes(x, 'Fixed-month behavior return roster')]
            pc[coord] = statistics.mean(selected) if selected else ''
        months = sorted([x['Return month'] for x in animal_rows if yes(x, 'Maintained calcium imaging') and x['Return month'] in ['3', '6', '9', '12']], key=int)
        other = any(yes(x, 'Maintained calcium imaging') and x['Return month'] not in ['3', '6', '9', '12'] for x in animal_rows)
        pub.cell(r, 7, ', '.join(months) + ('; other interval' if other else '') if months else ('Other intervals' if other else 'None'))
        pub.row_dimensions[r].height = 22
        for c in range(1, 8): pub.cell(r, c).border = Border(top=THIN if r > 4 else Side(), bottom=DARK if r == 21 else THIN, left=THIN if c > 1 else Side(), right=THIN if c < 7 else Side())
    style(pub, 'A4:G21', size=10, vertical='center')
    style(pub, 'A4:A21', size=10, bold=True, fill='F3F3F3', horizontal='center')
    style(pub, 'B4:F21', horizontal='center', number_format='0')
    style(pub, 'G4:G21', horizontal='left')
    for r, note in enumerate(notes, 23):
        pub.merge_cells(f'A{r}:G{r}'); pub.cell(r, 1, note)
        style(pub, f'A{r}:G{r}', size=10, fill=PALE, wrap=True, vertical='center')
        pub.row_dimensions[r].height = 58 if r == 26 else 42
    for c, width in zip('ABCDEFG', [11, 16, 16, 16, 16, 19, 28]): pub.column_dimensions[c].width = width
    pub.freeze_panes = 'A4'
    headers = ['Animal ID', 'Table return month (mo)', 'Previous nominal month (mo)', 'Preceding Day i session', 'Preceding Day i+1 session', 'Preceding Day i+1 date', 'Return Day i session', 'Return Day i date', 'Interval (d)', 'Fixed-month behavior return roster', 'Maintained calcium imaging', 'Pair key', 'Schedule note', 'Included in approximately 3-month summary']
    values(ledger, 1, [headers])
    for r, row in enumerate(rows, 2):
        vals = [row['Animal ID'], int(row['Return month']) if row['Return month'] in ['3', '6', '9', '12'] else None, previous_month(row['Previous nominal month']), row['Preceding Day i session'], row['Preceding Day i+1 session'], datetime.strptime(row['Preceding Day i+1 date'], '%Y-%m-%d'), row['Return Day i session'], datetime.strptime(row['Return Day i date'], '%Y-%m-%d'), f'=H{r}-F{r}', 'Yes' if yes(row, 'Fixed-month behavior return roster') else 'No', 'Yes' if yes(row, 'Maintained calcium imaging') else 'No', row['Pair key'], row['Schedule note'], 'No' if row['Pair key'] in [SHORT, DOUBLE] else 'Yes']
        values(ledger, r, [vals]); lc[f'I{r}'] = (vals[7] - vals[5]).days
        assert lc[f'I{r}'] == int(row['Interval (d)'])
        if row['Schedule note']: ledger.row_dimensions[r].height = 38
    style(ledger, 'A1:N1', size=9, color='FFFFFF', bold=True, fill=NAVY, wrap=True, horizontal='center', vertical='center')
    ledger.row_dimensions[1].height = 42
    style(ledger, 'A2:N59', size=9, vertical='center')
    for r in range(2, 60):
        for c in range(1, 15): ledger.cell(r, c).border = Border(top=THIN if r > 2 else Side(), bottom=THIN if r < 59 else Side())
    for address in ['A2:C59', 'F2:F59', 'H2:K59', 'N2:N59']: style(ledger, address, horizontal='center')
    for address in ['F2:F59', 'H2:H59']: style(ledger, address, number_format='yyyy-mm-dd')
    style(ledger, 'I2:I59', number_format='0'); style(ledger, 'M2:N59', wrap=True)
    for c, width in zip('ABCDEFGHIJKLMN', [11, 16, 16, 21, 21, 21, 21, 21, 13, 18, 18, 30, 48, 24]): ledger.column_dimensions[c].width = width
    ledger.freeze_panes = 'A2'
    table = Table(displayName='ComprehensiveIntervalLedger', ref='A1:N59')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True, showColumnStripes=False)
    ledger.add_table(table)
    summary.merge_cells('A1:G1'); summary['A1'] = 'Summary and provenance'
    style(summary, 'A1:G1', size=15, color='FFFFFF', bold=True, fill=NAVY, vertical='center')
    summary.row_dimensions[1].height = 30
    values(summary, 3, [['Cohort and interval check', 'Value']])
    labels = ['Displayed intervals', 'Displayed mice', 'Fixed-month behavior intervals', 'Fixed-month behavior mice', 'Calcium intervals (36 long + 1 short)', 'Behavior and calcium intervals', 'Behavior-only intervals', 'Additional calcium interval(s)', 'Approximately 3-month intervals', 'Mean (d)', 'Sample s.d. (d)', 'Minimum (d)', 'Maximum (d)', 'Mouse 618 double-gap exception (d)']
    formulas = ["=COUNT('Source Data'!$I$2:$I$59)", "=COUNTA(_xlfn.UNIQUE('Source Data'!$A$2:$A$59))", '=COUNTIF(\'Source Data\'!$J$2:$J$59,"Yes")', '=COUNTA(_xlfn.UNIQUE(_xlfn._xlws.FILTER(\'Source Data\'!$A$2:$A$59,\'Source Data\'!$J$2:$J$59="Yes")))', '=COUNTIF(\'Source Data\'!$K$2:$K$59,"Yes")', '=COUNTIFS(\'Source Data\'!$J$2:$J$59,"Yes",\'Source Data\'!$K$2:$K$59,"Yes")', '=COUNTIFS(\'Source Data\'!$J$2:$J$59,"Yes",\'Source Data\'!$K$2:$K$59,"No")', '=COUNTIFS(\'Source Data\'!$J$2:$J$59,"No",\'Source Data\'!$K$2:$K$59,"Yes")', '=COUNTIF(\'Source Data\'!$N$2:$N$59,"Yes")', '=AVERAGEIF(\'Source Data\'!$N$2:$N$59,"Yes",\'Source Data\'!$I$2:$I$59)', '=STDEV.S(_xlfn._xlws.FILTER(\'Source Data\'!$I$2:$I$59,\'Source Data\'!$N$2:$N$59="Yes"))', '=MINIFS(\'Source Data\'!$I$2:$I$59,\'Source Data\'!$N$2:$N$59,"Yes")', '=MAXIFS(\'Source Data\'!$I$2:$I$59,\'Source Data\'!$N$2:$N$59,"Yes")', '=AVERAGEIFS(\'Source Data\'!$I$2:$I$59,\'Source Data\'!$A$2:$A$59,"618",\'Source Data\'!$B$2:$B$59,12)']
    cached = [58, 18, 56, len({r['Animal ID'] for r in behavior}), 37, 35, 21, 2, 56, statistics.mean(days), statistics.stdev(days), min(days), max(days), 180]
    for r, (label, formula, result) in enumerate(zip(labels, formulas, cached), 4):
        summary.cell(r, 1, label); summary.cell(r, 2, formula); sc[f'B{r}'] = result
    style(summary, 'B4:B17', number_format='0'); style(summary, 'A8:A8', wrap=True)
    summary.row_dimensions[8].height = 30
    for r in range(3, 18):
        for c in [1, 2]: summary.cell(r, c).border = Border(top=THIN if r > 3 else Side(), bottom=THIN)
    style(summary, 'A3:B3', size=10, bold=True, fill=GRAY)
    for address, title in [('A20:G20', 'Completed manuscript sentence'), ('A26:G26', 'Explicit exception and exclusion'), ('A32:G32', 'Provenance')]:
        summary.merge_cells(address); summary[address.split(':')[0]] = title
        style(summary, address, size=10, color='FFFFFF', bold=True, fill=NAVY)
    summary.merge_cells('A21:G23'); summary['A21'] = SENTENCE
    style(summary, 'A21:G23', size=10, fill=PALE, wrap=True, vertical='center')
    for r in range(21, 24):
        summary.row_dimensions[r].height = 24
        for c in range(1, 8): summary.cell(r, c).border = Border(top=THIN if r == 21 else Side(), bottom=THIN if r == 23 else Side(), left=THIN if c == 1 else Side(), right=THIN if c == 7 else Side())
    values(summary, 27, [['618 exception', DOUBLE, '180 d', 'Shown under 12-month return; excluded from 90 +/- 4 d summary', None, None, None], ['1616 short interval', SHORT, '14 d', 'Shown under Other intervals; excluded from 90 +/- 4 d summary', None, None, None], ['Retained 1616 interval', RETAINED, '110 d', 'Shown under 6-month return; included in 90 +/- 4 d summary', None, None, None]])
    style(summary, 'A27:A29', size=9, bold=True); style(summary, 'A27:G29', fill='FFF2CC')
    summary.merge_cells('D28:G28'); style(summary, 'D28:G28', wrap=True)
    summary.row_dimensions[28].height = 36
    sha = lambda value: hashlib.sha256(value.encode('utf-8')).hexdigest()
    provenance = [['Comprehensive ledger', source_path.name, sha(text)], ['Ledger manifest', manifest_path.name, sha(manifest_text)]] + [['Upstream source', row['path'], row['sha256']] for row in manifest['source_files']]
    values(summary, 33, provenance); last = 32 + len(provenance)
    style(summary, f'A33:C{last}', size=9, color='555E68', wrap=True)
    for r in range(33, last + 1):
        summary.row_dimensions[r].height = 34
        for c in range(1, 4): summary.cell(r, c).border = Border(top=THIN if r > 33 else Side(), bottom=THIN if r < last else Side())
    style(summary, f'A33:A{last}', size=9, bold=True)
    for c, width in zip('ABCDEFG', [34, 31, 48, 18, 18, 18, 18]): summary.column_dimensions[c].width = width
    summary.freeze_panes = 'A2'
    output.mkdir(parents=True, exist_ok=True)
    path = output / 'Supplementary_Table_2_Long_Gap_Intervals.xlsx'
    book.save(path); cache_formulas(path, caches)
    check = load_workbook(path, data_only=True)
    assert check['Summary and provenance']['B12'].value == 56
    assert abs(check['Summary and provenance']['B14'].value - statistics.stdev(days)) < 1e-12
    assert not [c.coordinate for s in check for row in s for c in row if c.data_type == 'e']
    check.close()
    (output / 'Supplementary_Table_2_Long_Gap_Intervals.csv').write_bytes(source_path.read_bytes())
    (output / 'manuscript_sentence.tex').write_text(SENTENCE + '\n', encoding='utf-8')
    report = {'status': 'PASS', 'builder': 'openpyxl', 'displayed_intervals': 58, 'displayed_mice': 18, 'formula_count': sum(map(len, caches.values())), 'approximately_3_month_summary': {'n': 56, 'mean_days': statistics.mean(days), 'sample_sd_days': statistics.stdev(days), 'minimum_days': min(days), 'maximum_days': max(days)}, 'formula_cache_policy': 'Table-specific source-ledger calculation; Excel recalculates on opening', 'preview': 'Use the unchanged publication PDF renderer; openpyxl does not render worksheets.'}
    (output / 'MANIFEST.json').write_text(json.dumps(report, indent=2) + '\n', encoding='utf-8')
    return path


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('output_dir', type=Path, nargs='?', default=Path(__file__).resolve().parents[1] / 'build')
    parser.add_argument('source_dir', type=Path, nargs='?', default=Path(__file__).resolve().parents[1] / 'source_data')
    parser.add_argument('--notes-source', type=Path, default=Path(__file__).resolve().parents[1] / 'source_data')
    args = parser.parse_args()
    print(build(args.output_dir.resolve(), args.source_dir.resolve(), args.notes_source.resolve()))
