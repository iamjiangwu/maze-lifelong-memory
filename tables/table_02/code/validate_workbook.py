"""Validate Supplementary Table 2's formulas and retained-source cached values."""
import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


def validate(path):
    formula_book = load_workbook(path, data_only=False)
    book = load_workbook(path, data_only=True)
    publication = book['Supplementary Table 2']
    source = book['Source Data']
    summary = book['Summary and provenance']
    animals = {publication.cell(r, 1).value: r for r in range(4, 22)}
    assert publication.cell(animals['418'], 5).value == 90
    assert publication.cell(animals['618'], 5).value == 180
    r = animals['1616']
    assert [publication.cell(r, c).value for c in [3, 6, 7]] == [110, 14, '6; other interval']
    assert [summary[f'B{r}'].value for r in [4, 6, 8, 12]] == [58, 56, 37, 56]
    assert round(summary['B13'].value, 1) == 90.1 and round(summary['B14'].value, 1) == 3.6
    assert source.max_row == 59
    assert "excluded mouse 1616's 14-d interval" in publication['A25'].value
    assert "mouse 618's 180-d interval" in publication['A25'].value
    for r in range(2, 60):
        animal, month, *unused = [source.cell(r, c).value for c in range(1, 15)]
        interval = source.cell(r, 9).value
        assert interval == (source.cell(r, 8).value - source.cell(r, 6).value).days
        column = [3, 6, 9, 12].index(month) + 2 if month in [3, 6, 9, 12] else 6
        assert publication.cell(animals[animal], column).value == interval
    expected = {'1616day1->1616bday1': [None, 14, 'No'], '1616bday1->1616b-3day1': [6, 110, 'Yes'], '618-6day1->618-12day1': [12, 180, 'No']}
    for key, wanted in expected.items():
        matches = [r for r in range(2, 60) if source.cell(r, 12).value == key]
        assert len(matches) == 1
        assert [source.cell(matches[0], c).value for c in [2, 9, 14]] == wanted
    formulas = [(s.title, c.coordinate) for s in formula_book for row in s for c in row if c.data_type == 'f']
    assert len(formulas) == 162
    assert not [(s.title, c.coordinate) for s in book for row in s for c in row if c.data_type == 'e']
    for name, coordinate in formulas:
        assert book[name][coordinate].value is not None or name == 'Supplementary Table 2'
    formula_book.close(); book.close()
    return {'status': 'PASS', 'publication_rows': 18, 'source_rows': 58, 'formula_count': len(formulas), 'post_export_formula_errors': 0}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    print(json.dumps(validate(parser.parse_args().workbook), indent=2))
